"""Part A RC2: reflect the supplied vector logo kit. Original text edits keep run formatting."""
import sys, copy, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent)); sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[2] / "GEM_Brand_Assets_v1.0/02_build"))
from pptx import Presentation
from pptx.util import Inches
from pptx_helpers import find, set_text, replace_in_runs
import sync_logos as S
SRC = "GEM_V3_RC2/00_originals/RC2_pre_asset_sync/A-RC2_pre_asset_sync.pptx"
prs = Presentation(SRC); sl = prs.slides
def note_sub(slide, old, new):
    if not slide.has_notes_slide: return
    for p in slide.notes_slide.notes_text_frame.paragraphs:
        for r in p.runs:
            if old in r.text: r.text = r.text.replace(old, new)
print("logos swapped:", S.sync_deck(prs))
# slide 1
replace_in_runs(find(sl[0], "V3.0 RC2 · Working edition"), "logo shown from supplied raster, PENDING PRODUCTION MASTER", "logo from supplied vector kit, owner acceptance PENDING (VAL-02)")
note_sub(sl[0], "The logo is the supplied raster artwork, PENDING PRODUCTION MASTER (VAL-02).", "The logo is the supplied vector kit (SVG), used unchanged; owner acceptance as production master is PENDING (VAL-02).")
# slide 26 tags
for sh in sl[25].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "PENDING PRODUCTION MASTER": set_text(sh, "KIT · ACCEPTANCE PENDING")
note_sub(sl[25], "The logo is the supplied raster artwork, used unchanged. It is not a production master and is never trac", "The logo is the supplied vector kit (SVG), used unchanged. Acceptance as production master is pending (VAL-02); it is never trac")
# slide 27: logo set
s = sl[26]
set_text(find(s, "THE LOGO SET"), "THE LOGO SET · SUPPLIED VECTOR KIT")
set_text(find(s, "Five artworks exist"), "One mark, four artworks, vector throughout.")
cap = find(s, "Symbol, black on white")
ln_ref = [sh for sh in s.shapes if sh.name == "Shape 4"][0]._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ln")
for nm in ("Shape 6", "Shape 9"):
    sp = [sh for sh in s.shapes if sh.name == nm][0]._element.spPr; old = sp.find(ln_ref.tag); idx = list(sp).index(old); sp.remove(old); sp.insert(idx, copy.deepcopy(ln_ref))
for t in ("Stacked configuration", "Not supplied. PENDING PRODUCTION MASTER", "Micro, no-spark mark", "Required below 32px. Not supplied. PENDING"):
    e = find(s, t)._element; e.getparent().remove(e)
for left, kind, label in ((5.55, "stacked", "Stacked, black on white"), (7.88, "nospark", "No-spark, black on white")):
    S.add_kit_pic(s, kind, "black", Inches(left + 0.41), Inches(2.28), Inches(1.39), Inches(1.5), f"GEM {kind.replace('nospark','symbol without spark')} logo, black", 1400)
    c = copy.deepcopy(cap._element); s.shapes._spTree.append(c)
    from pptx.shapes.autoshape import Shape
    sh = Shape(c, s.shapes); sh.left = Inches(left + 0.01); sh.top = cap.top; sh.width = Inches(2.19); set_text(sh, label)
set_text(find(s, "Vector masters"), "Vector files")
set_text(find(s, "SVG, PDF, EPS, PNG."), "SVG, PDF, EPS, PNG. Acceptance PENDING (VAL-02)")
note_sub(s, "Raster logos must never be auto-traced.", "The vector kit replaces the earlier raster references; logos are never auto-traced or redrawn.")
# slide 28
set_text(find(sl[27], "Spark. Proportional measurements"), "Proportions follow the supplied vector kit. PENDING PRODUCTION MASTER (VAL-02).")
note_sub(sl[27], "Construction is described from the supplied artwork only.", "Construction is described from the supplied vector kit only.")
# slide 29
set_text(find(sl[28], "Below 32 px use"), "Below 32 px use the no-spark symbol (supplied kit). PENDING (VAL-02).")
# slide 30 note, 74 note, 78
note_sub(sl[29], "made from the supplied raster.", "made from the supplied vector kit.")
note_sub(sl[73], "Brand Masters are currently PENDING PRODUCTION MASTER.", "Brand Masters hold the supplied vector kit; owner acceptance as production master is PENDING (VAL-02).")
set_text(find(sl[77], "Production logo masters"), "Production logo masters · VAL-02 · VECTOR KIT RECEIVED, ACCEPTANCE PENDING")
prs.save("/tmp/claude-0/pd/A_new.pptx"); print("saved")
