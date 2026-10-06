"""Part A RC2 synchronization edits. Run from 05_logs. Edits 01_working/A-RC2.pptx in place
(the file is already a copy of the untouched original)."""
import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx_helpers import *

ROOT = pathlib.Path(__file__).resolve().parents[1]
F = ROOT / "01_working" / "A-RC2.pptx"
prs = Presentation(F)
S = lambda n: prs.slides[n - 1]
log = []
def L(slide, what): log.append((slide, what))

# ---------- 1 cover ----------
s = S(1)
set_text(find(s, "Working edition"), "V3.0 RC2 · Working edition · release not yet authorized (AC20) · logo shown from supplied raster, PENDING PRODUCTION MASTER · Document ID GEM-BG-V3.0-RC2 · issued 2026-10-06")
set_notes(s, "Cover. Edition: V3.0 Release Candidate 2 (RC2), a working edition synchronized with Part B and Part C RC2 and with the V3 Final Brand Approval Register. Brand Owner release authorization (AC20) is pending. The logo is the supplied raster artwork, PENDING PRODUCTION MASTER (VAL-02). Document control and change log: slide 79.")
L(1, "cover status line: V3.0 RC2, document ID, issue date")

# ---------- 3 authority ----------
s = S(3)
items = ["1 · V3 Final Brand Approval Register",
         "2 · Part A, Brand Guidelines: brand and identity",
         "3 · Part B, Digital Design System: digital implementation",
         "4 · Part C, Production Standards: physical production",
         "5 · Formal technical and legal standards, where they apply",
         "6 · V2 / V2.1 and V1: historical, superseded",
         "7 · Benchmarks and inspiration only"]
names = ["Text 3", "Text 5", "Text 7", "Text 9", "Text 11", "Text 14"]
for n, t in zip(names, items[:6]):
    set_text(by_name(s, n), t)
clone(s, by_name(s, "Text 14"), top=5.59, text=items[6], name="Text 27")
clone(s, by_name(s, "Shape 13"), top=5.96, name="Shape 28")
note = clone(s, by_name(s, "Text 24"), left=0.89, top=6.08, width=6.64, height=0.6, name="Text 29",
             text="Domain ownership: where A, B and C cover different implementation domains, the owning document governs that domain unless the Approval Register overrides it.")
set_font_size(note, 11)
set_text(by_name(s, "Text 24"), "Dielines, materials, tolerances, supplier proofs. Issued as RC2 · PENDING PRODUCTION VALIDATION.")
set_notes(s, "Authority order as set by the register (RF01, RF02, DS01): the register first; then each of the three V3.0 products in its own domain; formal technical and legal standards where they apply; V2/V2.1 and V1 are historical evidence only; benchmarks and inspiration never override a GEM rule. Part A cross-references B and C rather than repeating them. Part C exists as Release Candidate 2 and is a working edition pending production validation (AC17).")
L(3, "seven-step authority order + domain-ownership note; stale Part C tile removed")

# ---------- 4 status vocabulary ----------
s = S(4)
t = find(s, "Shape carries")
set_text(t, "Shape carries the meaning as well as the word: filled, solid, dashed, dotted. Placeholders read [PENDING], [TO BE VERIFIED] or [REQUIRES OWNER]. Concepts read CONCEPT / NOT PRODUCTION ARTWORK or CONCEPT / NOT PRODUCTION PHOTOGRAPHY. Further labels: PENDING PRODUCTION MASTER, PENDING LOCALIZATION APPROVAL, PROJECT-SPECIFIC / REQUIRES SITE VALIDATION. Edition states: WORKING EDITION, RELEASE CANDIDATE (RC), APPROVED. An edition becomes APPROVED only by Brand Owner authorization (AC20).")
move(t, height=1.55)
L(4, "declared the additional labels and the edition states")

# ---------- 23 secondary line weight ----------
replace_in_runs(find(S(23), "Secondary. Sentence case"), "Jost light", "Jost 400")
L(23, "Jost light → Jost 400")

# ---------- 33 colour limits ----------
set_text(find(S(33), "NEVER. 2.2"), "NEVER. 2.2 : 1 fails. Not for essential text, control boundaries or the logo (K05). Decorative use only (K07).")
L(33, "Beige-on-White wording aligned to K05/K07")

# ---------- 34 / 35 / 36 typography ----------
set_text(find(S(34), "Display and headings: 400"), "Display and headings: 400. Controlled labels and emphasis: 500. Two weights per composition at most. Display capitals carry .08 to .12em tracking (Part B tokens, CONDITIONAL · VAL-19).")
set_text(find(S(35), "LABEL · INTER 500"), "LABEL SMALL · INTER 400 · CAPITALS")
set_text(find(S(36), "DISPLAY CAPITALS"), "DISPLAY CAPITALS · .08–.12EM")
set_text(find(S(36), "Never track Arabic"), "Never track Arabic. The retired 0.3 to 0.8em rule is not used. Values are CONDITIONAL until optical QA in the approved templates (VAL-19). Full token values: Part B.")
L(34, "tracking range .08–.12em, CONDITIONAL VAL-19"); L(35, "label example mapped to Part B labelSmall 400"); L(36, "CONDITIONAL VAL-19 added")

# ---------- 38 bilingual composition ----------
t = find(S(38), "Both languages share")
set_text(t, "Both languages share one grid, one baseline rhythm and equal care. Neither is a footnote to the other. For Saudi/GCC bilingual wayfinding Arabic leads by default (S07, slide 65). Working strings only; not approved copy.")
move(t, height=0.75)
L(38, "cross-reference to the S07 wayfinding default")

# ---------- 6 / 41 / 46 geometry clear of text ----------
move(find(S(6), "Every arrival"), top=2.98)
s = S(41)
move(find(s, "OVERSIZED"), width=3.9)
t = find(s, "Let a form leave")
set_text(t, "Let a form leave the frame. A crop creates tension; the text keeps still. A form never intersects a headline or running text.")
move(t, width=3.9, height=1.0)
t = find(S(46), "Light, material")
move(t, width=4.9, height=0.75)
L(6, "headline moved below the ring"); L(41, "headline and body re-wrapped clear of the ring; rule wording"); L(46, "subtitle re-wrapped clear of the ring")

# ---------- 43 placement rule ----------
s = S(43)
set_text(find(s, "Forms are large, quiet"), "Forms are large, quiet and clear of the words.")
set_text(find(s, "Placement."), "Placement. Keep off the logo's clearspace and never behind or through text. Never repeat a form as a pattern.")
L(43, "placement rule: never behind or through text")

# ---------- 45 glyph wording ----------
set_text(find(S(45), "Glyphs help people"), "Glyphs help people find, choose and understand. They are geometric, reduced and single-colour: about a 2 px optical stroke on a defined optical box, drawn from the brand's circles and arcs (N15).")
L(45, "glyph description aligned to N15")

# ---------- 48 photography categories ----------
t = find(S(48), "REFERENCE · CONCEPT / NOT PRODUCTION PHOTOGRAPHY · PLAIN")
set_text(t, "REFERENCE · CONCEPT / NOT PRODUCTION PHOTOGRAPHY · PLAIN TILES AWAIT COMMISSIONED IMAGES. REQUIRED CATEGORIES (O11–O16): ARCHITECTURE, MATERIALS, THRESHOLDS, AMENITIES; DINING AND SERVICE WHERE RELEVANT. GUEST RITUAL, ROOM DETAILS AND HUMAN PRESENCE ARE SUPPLEMENTARY SUBJECTS.")
set_font_size(t, 10); move(t, top=6.47, width=10.2, height=0.6)
L(48, "subjects mapped to register categories O11–O16")

# ---------- 50 / 51 status labels ----------
set_text(find(S(50), "REFERENCE · CONCEPT"), "REFERENCE · CONCEPT / NOT PRODUCTION PHOTOGRAPHY · RATIOS CONDITIONAL (W07)")
set_text(find(S(51), "About 1.5 seconds"), "About 1.5 seconds, ease-out, once. Reduced motion: show the final frame at rest, no animation. Reveal timings CONDITIONAL; motion masters PENDING VALIDATION (VAL-14). Motion tokens: Part B.")
L(50, "CONDITIONAL (W07)"); L(51, "CONDITIONAL / VAL-14")

# ---------- 53 radius ----------
set_text(find(S(53), "Square edges"), "Near-square edges (control radius per Part B), hairline fields, one clear action.")
L(53, "radius wording aligned to Part B")

# ---------- 60 tiers ----------
s = S(60)
set_text(find(s, "The consistent GEM range"), "The consistent GEM range. Name [REQUIRES OWNER]. Scope [PENDING]")
t = find(s, "Same logo, same palette")
set_text(t, "Same logo, same palette, same type in every tier. All tiers are CONDITIONAL until commercial scope closes (AC03, VAL-01). Property Edition and Bespoke programmes are named in the register (E14); the name Core Collection is [REQUIRES OWNER].")
move(t, height=0.75)
L(60, "tier status CONDITIONAL AC03/VAL-01; Core Collection name [REQUIRES OWNER]")

# ---------- 61 six levels (T05) ----------
s = S(61)
remove(by_name(s, "Text 13")); remove(by_name(s, "Text 14")); remove(by_name(s, "Shape 12"))
set_text(by_name(s, "Text 11"), "Product / function")
move(by_name(s, "Text 11"), width=2.6)
for n in ["Shape 15", "Text 16", "Text 17", "Shape 18", "Text 19", "Text 20", "Shape 21", "Shape 22", "Text 23", "Text 24"]:
    shift(by_name(s, n), -0.52)
set_text(by_name(s, "Text 16"), "04"); set_text(by_name(s, "Text 19"), "05"); set_text(by_name(s, "Text 23"), "06")
set_text(by_name(s, "Text 17"), "Quantity / size"); move(by_name(s, "Text 17"), width=1.8)
set_text(by_name(s, "Text 20"), "Instructions / market information"); move(by_name(s, "Text 20"), width=3.4)
set_text(find(s, "The brand is seen first"), "The brand is seen first, then the tier, then the product and its function, then quantity, information and the SKU. Six levels, as approved (T05). Never reverse the order.")
set_notes(s, "Type size steps down with each level. Six levels per register decision T05: GEM → collection/tier → product/function → quantity/size → instructions/required market information → SKU/artwork version. Detailed layout, dielines and production controls belong to Part C.")
L(61, "hierarchy aligned to the six T05 levels")

# ---------- 65 wayfinding ----------
s = S(65)
t = find(s, "One direction, one message")
set_text(t, "One direction, one message per sign. Read left to right in English, right to left in Arabic. For Saudi/GCC bilingual wayfinding Arabic leads by default: Arabic first and right, English secondary and left, with equivalent information and legibility (S07). Geometry is shared with the brand.")
move(t, top=5.38, height=0.85)
t = find(s, "PROJECT-SPECIFIC / REQUIRES SITE VALIDATION")
set_text(t, "PROJECT-SPECIFIC / REQUIRES SITE VALIDATION · ARABIC: PENDING LOCALIZATION APPROVAL · EXTERNAL PROPERTY SIGNAGE: DEFERRED (S02)")
move(t, top=6.3)
set_notes(s, "Fabrication, sizes, mounting and materials depend on each property and are not specified here (Part C). Bilingual order for Saudi/GCC wayfinding follows register decision S07: Arabic first and right, English secondary and left; adapt only where regulatory or operational context requires, as a documented project-specific deviation. External property signage is deferred (S02). Arabic text is a placeholder pending localization approval.")
L(65, "S07 default stated once; S02 deferral noted")

# ---------- 68 / 69 templates ----------
set_text(find(S(68), "CONCEPT / NOT PRODUCTION ARTWORK · ARABIC"), "CONCEPT / NOT PRODUCTION ARTWORK · ARABIC: PENDING LOCALIZATION APPROVAL · TEMPLATES: OPEN DELIVERABLE (AB10, W01, W13)")
set_text(find(S(69), "One form, one line"), "One form, one line, one spark at most. No loud colour, no stickers, no gradients. Social starter set: OPEN DELIVERABLE (W07), not yet produced.")
L(68, "template deliverable gate"); L(69, "social template deliverable gate")

# ---------- 71 accessibility notes ----------
set_notes(S(71), "Accessibility implementation QA remains an open gate (VAL-08). This deck uses readable sizes and AA contrast pairings; every image carries alt text and purely decorative geometry is marked decorative (V02, V03). Native document QA (VAL-15) is still to run.")
L(71, "alt-text claim made true (see alt text pass)")

# ---------- 72 twelve roles ----------
s = S(72)
set_text(find(s, "EIGHT ROLES"), "TWELVE ROLES KEEP GEM EXACT.")
roles = ["Brand Owner", "Design Custodian", "Commercial Lead", "Product / Amenities Lead",
         "Procurement / Supplier QA", "Digital Design Lead", "Engineering Lead", "Arabic / Localization Lead",
         "Accessibility QA", "Legal / IP Counsel", "Presentation / Document QA", "Asset Librarian / Governance PM"]
tiles = [("Shape 2", "Text 3", "Text 4"), ("Shape 5", "Text 6", "Text 7"), ("Shape 8", "Text 9", "Text 10"), ("Shape 11", "Text 12", "Text 13"),
         ("Shape 14", "Text 15", "Text 16"), ("Shape 17", "Text 18", "Text 19"), ("Shape 20", "Text 21", "Text 22"), ("Shape 23", "Text 24", "Text 25")]
xs = [0.90, 3.82, 6.74, 9.67]; ys = [1.95, 3.30, 4.65]; th = 1.2
objs = [(by_name(s, a), by_name(s, b), by_name(s, c)) for a, b, c in tiles]
srcs = objs[0]
for i in range(8, 12):
    objs.append(tuple(clone(s, src, name=f"{src.name}-{i}") for src in srcs))
for i, (box, name, status) in enumerate(objs):
    x, y = xs[i % 4], ys[i // 4]
    move(box, left=x, top=y, height=th)
    move(name, left=x + 0.17, top=y + 0.15, height=0.6); set_text(name, roles[i])
    move(status, left=x + 0.17, top=y + th - 0.4); set_text(status, "[REQUIRES OWNER]")
t = find(s, "Roles are defined")
set_text(t, "Roles follow the V3 Approval Register (Owners & Governance). Named holders are not yet assigned (AC19). Release authorization (AC20) rests with the Brand Owner. ™ in running text and on packs: [REQUIRES OWNER · Legal / IP Counsel].")
move(t, top=6.03, height=0.6)
set_notes(s, "The twelve roles and their responsibilities are those of the register's Owners & Governance sheet; Part B and Part C cite this list and add only their domain duties. No people are invented. Detailed RACI is intentionally omitted from Part A.")
L(72, "twelve canonical register roles; ™ placeholder")

# ---------- 75 file naming (X12) ----------
s = S(75)
set_text(find(s, "GEM_[Category]"), "GEM_[STREAM]_[ASSET]_[VARIANT]_vX.Y_YYYYMMDD.ext")
set_text(find(s, "Example:"), "Example: GEM_Logo_Horizontal_Black_v3.0_20261006.svg · convention X12 of the V3 Approval Register")
set_text(find(s, "No spaces."), "No spaces. Short, predictable names. The stream field follows the asset library folder.")
set_text(find(s, 'No "final"'), 'No "final", "new" or "latest". The date field (YYYYMMDD) is the issue date; the version number carries asset history.')
set_text(find(s, "Approval date, checksum"), "Approval, checksum, owner, source and release status live in the asset manifest. Proposed, CONDITIONAL [REQUIRES OWNER]: asset ID GEM-[CATEGORY]-[NNN] and SHA-256 checksums (Part C, Appendix N).")
set_notes(s, "Convention X12 of the register. Earlier working editions omitted the date field; the register wins. Asset-ID format and checksum algorithm are proposals pending owner confirmation.")
L(75, "X12 naming adopted; asset ID / SHA-256 recorded as proposals")

# ---------- 78 open items with VAL IDs ----------
s = S(78)
vals = {"Text 3": ("Production logo masters ", "· VAL-02 · PENDING PRODUCTION MASTER"),
        "Text 5": ("Packaging pilot proof ", "· VAL-09 · NOT STARTED"),
        "Text 7": ("Trademark and artwork ownership ", "· VAL-03, VAL-04 · [TO BE VERIFIED]"),
        "Text 9": ("Production standards, Part C RC2 ", "· ISSUED · AC17, VAL-10, VAL-11 OPEN"),
        "Text 11": ("Font build and licence ", "· VAL-05 · PENDING VALIDATION"),
        "Text 13": ("Accessibility implementation QA ", "· VAL-08 · IN PROGRESS"),
        "Text 15": ("Image rights and commissioned photography ", "· VAL-06 · PENDING"),
        "Text 17": ("Governance holders ", "· AC19 · [REQUIRES OWNER]"),
        "Text 20": ("Arabic native QA and localization ", "· VAL-07, AC10 · PENDING VALIDATION"),
        "Text 23": ("Release authorization ", "· AC20 · [REQUIRES OWNER]")}
for n, (a, b) in vals.items():
    set_runs(by_name(s, n), [a, b])
more = clone(s, by_name(s, "Text 24"), left=0.89, top=6.47, width=11.9, height=0.3, name="Text 26",
             text="Also open: VAL-13, VAL-14, VAL-15, VAL-16, VAL-17, VAL-19, VAL-20, VAL-21 and the template deliverables (AB10, AB11, W01–W14). IDs follow the V3 Approval Register.")
set_font_size(more, 11)
set_notes(s, "Summarized by theme with the register's validation IDs (VAL-01 to VAL-21; there is no VAL-18 in the register) and the release checkpoints AC17, AC19 and AC20. Editing this manual closes none of them.")
L(78, "open items carry register VAL / AC IDs; Part C row current")

# ---------- 79 document control (duplicated from 78) ----------
s = S(79)
set_text(by_name(s, "Text 0"), "DOCUMENT CONTROL & CHANGE LOG")
set_text(by_name(s, "Text 1"), "ONE VERSION. ONE DATE. ONE OWNER.")
ctrl = {"Text 3": ("Document ID ", "· GEM-BG-V3.0-RC2"),
        "Text 5": ("Title ", "· GEM™ Brand Guidelines V3.0 — Part A"),
        "Text 7": ("Version ", "· V3.0 RC2 (Release Candidate 2)"),
        "Text 9": ("Edition state ", "· WORKING EDITION · NOT RELEASED"),
        "Text 11": ("Issue date ", "· 2026-10-06"),
        "Text 13": ("Owner ", "· Brand Owner [REQUIRES OWNER]"),
        "Text 15": ("Approver ", "· Brand Owner · AC20 PENDING"),
        "Text 17": ("Release status ", "· RC2 synchronized · evidence gates open"),
        "Text 20": ("Supersedes ", "· V3.0 working edition (undated, pre-RC2)"),
        "Text 23": ("Related ", "· GEM-DDS-V3.0-RC2 (Part B) · GEM-PS-V3.0-RC2 (Part C) · V3 Approval Register")}
for n, (a, b) in ctrl.items():
    set_runs(by_name(s, n), [a, b])
cl = clone(s, by_name(s, "Text 24"), left=0.89, top=6.42, width=11.9, height=0.36, name="Text 26",
           text="Change log · RC2, 2026-10-06: synchronized with Part B and Part C RC2 and the V3 Approval Register — authority order, twelve governance roles, VAL IDs, S07 wayfinding default, six-level hierarchy (T05), X12 file naming, status and edition vocabulary, alt text. Full record: GEM_V3_RC2_Release_Notes.")
set_font_size(cl, 10)
set_text(by_name(s, "Text 24"), "GEM™ BRAND GUIDELINES V3.0 · PART A · STEWARDSHIP · DOCUMENT CONTROL")
set_text(by_name(s, "Text 25"), "79")
set_notes(s, "Document control per register A11 and the change-log appendix (Z24, AB19). Owner and approver names are recorded by the Brand Owner before AC20. Part B's document ID is applied through its patch specification until its source is regenerated.")
L(79, "new document-control and change-log slide")

# ---------- 80 closing ----------
s = S(80)
set_text(find(s, "Brand Guidelines V3.0 · Part A"), "Brand Guidelines V3.0 RC2 · Part A · Working edition · not released")
set_notes(s, "Closing. Part B is the Digital Design System V3.0 (RC2 patch specification issued; PDF regeneration pending its source). Part C, Production Standards V3.0, is issued as Release Candidate 2, a working edition pending production validation.")
L(80, "closing edition line; stale note removed")

# ---------- footers ----------
n = replace_everywhere(prs, "GEM™ BRAND GUIDELINES V3.0 · PART A", "GEM™ BRAND GUIDELINES V3.0 RC2 · PART A")
L(0, f"footer version string on {n} shapes")

# ---------- alt text ----------
DEC = "Decorative background geometry: cropped GEM ring and quarter forms."
ALT = {
 (1,"Image 0"): (DEC, True), (1,"Image 1"): ("GEM primary horizontal signature, Beige on Ink, supplied raster, PENDING PRODUCTION MASTER", False),
 (5,"Image 0"): (DEC, True), (6,"Image 0"): ("Beige field with one cropped ring line, decorative.", True), (9,"Image 0"): (DEC, True),
 (18,"Image 0"): (DEC, True), (25,"Image 0"): (DEC, True), (52,"Image 0"): (DEC, True), (59,"Image 0"): (DEC, True), (64,"Image 0"): (DEC, True), (70,"Image 0"): (DEC, True),
 (8,"Image 0"): ("Rendered reference study: a ring pressed into a warm uncoated surface. Concept, not production photography.", False),
 (10,"Image 0"): ("Schematic of the four device forms: ring, opening arc, quarter mass and four-point spark. Not the master artwork.", False),
 (11,"Image 0"): ("Rendered reference study: a lit doorway threshold in a dark field. Concept, not production photography.", False),
 (26,"Image 0"): ("Primary horizontal signature, Beige on Ink, supplied raster.", False), (26,"Image 1"): ("Monochrome signature, Black on White, supplied raster.", False), (26,"Image 2"): ("Reversed signature, White on Ink, supplied raster.", False),
 (27,"Image 0"): ("Standalone symbol, Beige on Ink, supplied raster.", False), (27,"Image 1"): ("Standalone symbol, Black on White, supplied raster.", False),
 (28,"Image 0"): ("GEM symbol enlarged with construction callouts: G-ring, opening, quarter mass and spark, described from the supplied raster.", False), (28,"Image 1"): (DEC, True),
 (29,"Image 0"): ("Primary signature inside the dashed one-u clearspace boundary.", False), (29,"Image 1"): ("Primary lockup at its 120 px minimum digital width.", False), (29,"Image 2"): ("Symbol at its 32 px minimum digital size.", False),
 (30,"Image 0"): ("Logo misuse example: stretched.", False), (30,"Image 1"): ("Logo misuse example: rotated.", False), (30,"Image 2"): ("Logo misuse example: recoloured outside the palette.", False), (30,"Image 3"): ("Logo misuse example: drop shadow added.", False), (30,"Image 4"): ("Logo misuse example: placed on a gradient background.", False), (30,"Image 5"): ("Logo misuse example: outlined and contained.", False), (30,"Image 6"): ("Logo misuse example: cropped.", False),
 (31,"Image 0"): ("GEM signature beside a placeholder partner mark, equal-partner co-branding concept.", False), (31,"Image 1"): ("GEM signature in the supplier-relationship concept, placed after 'Supplied by'.", False),
 (40,"Image 0"): ("Composition schematic: margin, edge, crop and open field. Not a template.", False),
 (41,"Image 0"): ("Oversized ring cropped by the right edge of a dark field, demonstrating scale and tension.", False),
 (42,"Image 0"): ("Illustrative ring drawing.", False), (42,"Image 1"): ("Illustrative arc drawing.", False), (42,"Image 2"): ("Illustrative quarter drawing.", False), (42,"Image 3"): ("Illustrative four-point spark drawing.", False),
 (43,"Image 0"): ("Diagram of outline forms placed large in the open field, off the logo clearspace.", False),
 (44,"Image 0"): ("Illustrative four-point spark drawing, not the master artwork.", False),
 (45,"Image 0"): ("Illustrative direction glyph: arrow.", False), (45,"Image 1"): ("Illustrative confirm glyph: check in a circle.", False), (45,"Image 2"): ("Illustrative information glyph: i in a circle.", False),
 (46,"Image 0"): (DEC, True),
 (47,"Image 0"): ("Rendered reference study: pressed ring in warm single-source light. Concept, not production photography.", False),
 (48,"Image 0"): ("Reference study: threshold.", False), (48,"Image 1"): ("Reference study: material close crop.", False), (48,"Image 2"): ("Reference study: dining table, overhead.", False), (48,"Image 3"): ("Reference study: room detail, window light.", False),
 (49,"Image 0"): ("Do: one source of light, quiet tones, room to breathe.", False), (49,"Image 1"): ("Don't: oversaturated with artificial glow.", False), (49,"Image 2"): ("Don't: text placed over the subject.", False), (49,"Image 3"): ("Don't: trendy filter and off-palette tint.", False),
 (50,"Image 0"): ("Threshold study cropped 3:2.", False), (50,"Image 1"): ("Threshold study cropped 4:5.", False), (50,"Image 2"): ("Threshold study cropped 16:9.", False), (50,"Image 3"): ("Threshold study cropped 1:1.", False), (50,"Image 4"): ("Threshold study cropped 9:16.", False),
 (51,"Image 0"): ("Storyboard frame 1: the ring draws itself.", False), (51,"Image 1"): ("Storyboard frame 2: the spark appears once.", False), (51,"Image 2"): ("Storyboard frame 3: the quarter fills.", False),
 (53,"Image 0"): ("GEM signature in the homepage concept header.", False), (53,"Image 1"): ("Homepage concept: cropped ring behind a quiet arrival headline. Concept, not production artwork.", False),
 (55,"Image 0"): ("Guest journey timeline with six moments.", False),
 (56,"Image 0"): ("Key card concept with a cropped ring. Concept, not production artwork.", False), (56,"Image 1"): ("GEM signature on the key card concept.", False),
 (57,"Image 0"): ("Rendered reference study: window light on a wall. Concept, not production photography.", False),
 (58,"Image 0"): ("Rendered reference study: table setting from above, few objects. Concept, not production photography.", False),
 (62,"Image 0"): ("GEM symbol on the Shampoo concept pack.", False), (62,"Image 1"): ("GEM symbol on the Conditioner concept pack.", False), (62,"Image 2"): ("GEM symbol on the Body Wash concept pack.", False), (62,"Image 3"): ("GEM symbol on the Body Lotion concept pack.", False), (62,"Image 4"): ("GEM symbol on the Hand Wash concept pack.", False),
 (63,"Image 0"): ("Rendered reference study: pressed ring on one surface. Concept, not production photography.", False),
 (65,"Image 0"): ("Illustrative direction arrow glyph.", False),
 (66,"Image 0"): ("Illustrative information glyph, English sign concept.", False), (66,"Image 1"): ("Illustrative information glyph, Arabic sign concept.", False),
 (67,"Image 0"): ("GEM signature on the letterhead concept.", False),
 (68,"Image 0"): ("Miniature bar marks on the data slide concept.", False), (68,"Image 1"): ("Image-led slide miniature using the threshold study.", False),
 (69,"Image 0"): ("1:1 social concept with a cropped ring.", False), (69,"Image 1"): ("9:16 social concept with a quarter form.", False),
 (77,"Image 0"): ("GEM primary signature, quick reference.", False), (77,"Image 1"): ("Ring, arc, quarter and spark drawings, quick reference.", False),
 (80,"Image 0"): (DEC, True), (80,"Image 1"): ("GEM primary horizontal signature, Beige on Ink, supplied raster.", False),
}
done = 0
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.shape_type == 13:
            d, dec = ALT[(i, sh.name)]
            set_alt(sh, d, dec); done += 1
L(0, f"alt text on {done} pictures")

prs.save(F)
for sl, what in log: print(f"A {sl:>2}: {what}")
