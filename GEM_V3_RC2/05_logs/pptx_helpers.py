"""Shared helpers for RC2 edits. Edits run text in place so font, size, tracking
and colour are preserved; never uses text_frame.text = ... on styled shapes."""
import copy
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

DECOR_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
ADEC_NS = "http://schemas.microsoft.com/office/drawing/2017/decorative"


def find(slide, startswith, nth=0):
    hits = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip().startswith(startswith)]
    if len(hits) <= nth:
        raise KeyError(f"slide: no shape starting with {startswith!r} (#{nth})")
    return hits[nth]


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def set_text(shape, text):
    """Replace all text with one run that keeps the first run's formatting."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    runs = p0.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        r = p0.add_run(); r.text = text
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    return shape


def set_runs(shape, texts):
    """Assign text to the existing runs of the first paragraph, in order."""
    runs = shape.text_frame.paragraphs[0].runs
    assert len(runs) >= len(texts), (shape.name, len(runs), texts)
    for r, t in zip(runs, texts):
        r.text = t
    for r in runs[len(texts):]:
        r._r.getparent().remove(r._r)
    return shape


def replace_in_runs(shape, old, new):
    n = 0
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new); n += 1
    return n


def set_font_size(shape, pt):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(pt)


def move(shape, left=None, top=None, width=None, height=None):
    if left is not None: shape.left = Inches(left)
    if top is not None: shape.top = Inches(top)
    if width is not None: shape.width = Inches(width)
    if height is not None: shape.height = Inches(height)
    return shape


def shift(shape, dy):
    shape.top = shape.top + Inches(dy)


def clone(slide, src, left=None, top=None, width=None, height=None, text=None, name=None):
    el = copy.deepcopy(src._element)
    slide.shapes._spTree.append(el)
    new = slide.shapes[-1]
    # unique id
    ids = [s.shape_id for s in slide.shapes if s is not new]
    new._element.xpath("./p:nvSpPr/p:cNvPr")[0].set("id", str(max(ids) + 1))
    if name: new._element.xpath("./p:nvSpPr/p:cNvPr")[0].set("name", name)
    move(new, left, top, width, height)
    if text is not None: set_text(new, text)
    return new


def remove(shape):
    shape._element.getparent().remove(shape._element)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def set_alt(pic, descr, decorative=False):
    cNvPr = pic._element.xpath("./p:nvPicPr/p:cNvPr")[0]
    cNvPr.set("descr", descr)
    cNvPr.set("title", descr[:60])
    if decorative:
        extLst = cNvPr.find(qn("a:extLst"))
        if extLst is None:
            extLst = etree.SubElement(cNvPr, qn("a:extLst"))
        ext = etree.SubElement(extLst, qn("a:ext")); ext.set("uri", DECOR_URI)
        d = etree.SubElement(ext, "{%s}decorative" % ADEC_NS); d.set("val", "1")


def replace_everywhere(prs, old, new):
    n = 0
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                n += replace_in_runs(sh, old, new)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        for p in c.text_frame.paragraphs:
                            for r in p.runs:
                                if old in r.text: r.text = r.text.replace(old, new); n += 1
    return n


def set_cell(cell, text):
    p0 = cell.text_frame.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for r in p0.runs[1:]: r._r.getparent().remove(r._r)
    else:
        p0.add_run().text = text
    for p in cell.text_frame.paragraphs[1:]: p._p.getparent().remove(p._p)


def defit(shape):
    """Remove shrink-on-overflow / auto-grow so text renders at its set size."""
    bodyPr = shape.text_frame._txBody.find(qn("a:bodyPr"))
    for tag in ("a:normAutofit", "a:spAutoFit"):
        el = bodyPr.find(qn(tag))
        if el is not None:
            bodyPr.remove(el)
    return shape
