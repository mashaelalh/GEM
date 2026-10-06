"""Parts B and C RC2: reflect the supplied vector logo kit. Gates stay open; nothing is marked approved."""
import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent)); sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[2] / "GEM_Brand_Assets_v1.0/02_build"))
from pptx import Presentation
from pptx_helpers import find, set_text, replace_in_runs, set_cell
import sync_logos as S
def note_sub(slide, old, new):
    if not slide.has_notes_slide: return 0
    n = 0
    for p in slide.notes_slide.notes_text_frame.paragraphs:
        for r in p.runs:
            if old in r.text: r.text = r.text.replace(old, new); n += 1
    return n
# ---------------- Part C
C = Presentation("GEM_V3_RC2/00_originals/RC2_pre_asset_sync/C-RC2_pre_asset_sync.pptx"); sl = C.slides
print("C logos swapped:", S.sync_deck(C))
replace_in_runs(find(sl[0], "Release conditions incomplete"), "logo shown from supplied raster, PENDING PRODUCTION MASTER", "logo from supplied vector kit, acceptance PENDING (VAL-02)")
assert note_sub(sl[0], "Logo is the supplied raster artwork, not a production master (VAL-02).", "Logo is the supplied vector kit (SVG), used unchanged; its acceptance as production master is pending (VAL-02).")
set_text(find(sl[2], "Master file does not yet exist."), "Master file not yet accepted.")
t = [sh for sh in sl[6].shapes if sh.has_table][0].table
KIT = "SVG, PDF, EPS, PNG supplied"; ST = "KIT RECEIVED · ACCEPTANCE PENDING"
for r in t.rows:
    a = r.cells[0].text
    if a in ("Primary horizontal signature", "Monochrome (Black)", "Reversed (White)", "Standalone Black symbol", "Standalone Beige symbol", "Standalone White symbol", "Stacked configuration", "Micro / no-spark mark"):
        set_cell(r.cells[1], KIT); set_cell(r.cells[2], ST)
    if a.startswith("Favicon"):
        set_cell(r.cells[1], "Icon set supplied (kit)"); set_cell(r.cells[2], ST)
set_text(find(sl[6], "PNG is for raster use only."), "Vector kit supplied: SVG, PDF, EPS, PNG (see Part B). Controlled names follow X12 at first manifest (VAL-16); files are renamed, never redrawn. Acceptance as production master PENDING (VAL-02).")
note_sub(sl[6], "Raster files supplied with Part A are references for layout, not production masters.", "The vector kit supplied with this release replaces the earlier raster references; it is not a production master until the Brand Owner accepts it (VAL-02).")
set_text(find(sl[9], "Micro / no-spark mark: PENDING"), "No-spark symbol: supplied in kit · ACCEPTANCE PENDING")
for old, new in (("Production vector logo masters · VAL-02", "Logo masters, kit received · VAL-02"), ("Micro / no-spark logo · VAL-02", "No-spark logo, kit received · VAL-02"), ("Stacked logo · VAL-02", "Stacked logo, kit received · VAL-02"), ("White symbol · VAL-02", "White symbol, kit received · VAL-02")):
    set_text(find(sl[74], old), new)
C.save("/tmp/claude-0/pd/C_new.pptx")
# ---------------- Part B
B = Presentation("GEM_V3_RC2/00_originals/RC2_pre_asset_sync/B-RC2_pre_asset_sync.pptx"); sl = B.slides
print("B logos swapped:", S.sync_deck(B))
assert note_sub(sl[0], "Logo artwork is not reproduced here: production masters are PENDING (VAL-02).", "Logo artwork is not reproduced here; the supplied vector kit is received and its acceptance as production master is PENDING (VAL-02).")
tb = [sh for sh in sl[30].shapes if sh.has_table][0].table
for r in tb.rows:
    if r.cells[0].text == "Production": set_cell(r.cells[2], "Vector logo kit received, acceptance pending; no supplier data; Part C issued as RC2, physical evidence open (VAL-02, VAL-09 to VAL-11, AC17)")
set_text(find(sl[34], "ASSETS · PENDING"), "ASSETS · VECTOR KIT RECEIVED · ACCEPTANCE PENDING")
tl = [sh for sh in sl[34].shapes if sh.has_table][0].table
rows = [("gem-horizontal-beige.svg", "BEIGE", "Signature lockup on INK"), ("gem-horizontal-black.svg", "BLACK", "Monochrome lockup on white"), ("gem-horizontal-white.svg", "WHITE", "Reversed on INK or black only"), ("gem-symbol-beige.svg", "BEIGE", "Standalone symbol on INK"), ("gem-symbol-black.svg", "BLACK", "Standalone symbol on white")]
for r, v in zip(list(tl.rows)[1:], rows):
    for c, x in zip(r.cells, v): set_cell(c, x)
set_text(find(sl[34], "Rules (H10"), "Rules (H10–H15). Lockup min 120px; symbol min 32px; below that the no-spark symbol (H12; supplied in the kit). Clearspace u = G-ring radius, min 1u, preferred 2u. Never mirror (M08) or recreate in type (L16).")
set_text(find(sl[34], "Required, not supplied."), "Also in the kit: ink and white colourways, stacked lockup, no-spark symbol, clearspace versions, PDF, EPS and PNG, favicon and app icons (H16–H20). Renamed to GEM_[STREAM]_[ASSET]_[VARIANT]_vX.Y_YYYYMMDD.ext (X12) at first manifest (VAL-16), never redrawn. Embroidery and deboss masters not supplied. Acceptance PENDING (VAL-02).")
set_text(find(sl[34], "GEM™ Digital Design System V3.0 RC2 · working specification · 35"), "GEM™ Digital Design System V3.0 RC2 · working specification · 35 · Logos are the supplied vector kit; never redrawn or auto-traced (DS05).")
B.save("/tmp/claude-0/pd/B_new.pptx"); print("saved")
