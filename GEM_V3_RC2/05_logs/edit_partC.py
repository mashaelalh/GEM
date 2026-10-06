"""Part C RC2 synchronization edits on 01_working/C-RC2.pptx (a copy of the untouched original)."""
import sys, pathlib, copy
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from pptx_helpers import *

ROOT = pathlib.Path(__file__).resolve().parents[1]
F = ROOT / "01_working" / "C-RC2.pptx"
prs = Presentation(F)
S = lambda n: prs.slides[n - 1]
log = []
def L(n, w): log.append((n, w))

def table_of(slide):
    return [sh for sh in slide.shapes if sh.shape_type == 19][0]

def add_row(gf, texts):
    tbl = gf.table._tbl
    trs = tbl.findall(qn("a:tr"))
    new = copy.deepcopy(trs[-1]); tbl.append(new)
    row = gf.table.rows[len(gf.table.rows) - 1]
    for c, t in zip(row.cells, texts): set_cell(c, t)
    gf.height = gf.height + int(trs[-1].get("h"))
    return row

def del_row(gf, idx):
    tr = gf.table._tbl.findall(qn("a:tr"))[idx]
    gf.height = gf.height - int(tr.get("h")); tr.getparent().remove(tr)

def cell(gf, r, c): return gf.table.cell(r, c)

# ---------- 1 cover ----------
s = S(1)
set_text(find(s, "RELEASE CANDIDATE 1"), "RELEASE CANDIDATE 2")
t = find(s, "Release conditions incomplete"); defit(t)
set_text(t, "Release conditions incomplete · release authorization [REQUIRES OWNER] (AC20) · logo shown from supplied raster, PENDING PRODUCTION MASTER · Document ID GEM-PS-V3.0-RC2 · V3.0 RC2 · issued 2026-10-06")
set_notes(s, "Cover. Part C governs physical production and supplier release. Edition: V3.0 Release Candidate 2 (RC2), a working edition synchronized with Part A and Part B RC2 and the V3 Final Brand Approval Register. It may be labelled APPROVED PRODUCTION STANDARD only when the six release gates are met and the Brand Owner authorizes release (AC20). Logo is the supplied raster artwork, not a production master (VAL-02). Document control and change log: slide 77.")
L(1, "RC2, document ID, issue date")

# ---------- 2 authority ----------
s = S(2)
t = find(s, "1 V3 approval register"); defit(t)
set_text(t, "1 V3 Final Brand Approval Register · 2 Part A, brand and identity · 3 Part B, digital implementation · 4 Part C, physical production · 5 formal technical and legal standards where they apply · 6 V2 / V2.1 and V1, historical and superseded · 7 benchmarks and inspiration only. Domain ownership: the owning document governs its domain unless the register overrides it.")
move(t, height=1.0)
set_notes(s, "Authority order shared with Part A slide 3 and Part B (RF01, RF02, DS01). Part A governs what GEM should look and feel like, Part B digital implementation, Part C physical production and supplier release; where domains differ, the owning document governs unless the register overrides it. Approved brand decisions in the register are not reopened here.")
L(2, "seven-step authority order with domain-ownership note")

# ---------- 3 status vocabulary ----------
s = S(3)
t = find(s, "Placeholders:"); defit(t)
set_text(t, "Placeholders: [PENDING] [TO BE VERIFIED] [TO BE VERIFIED BY PROCESS TEST] [REQUIRES OWNER] [REQUIRES SUPPLIER] [REQUIRES SUPPLIER / OWNER] [ENTER] [PENDING PRODUCTION MASTER] [PENDING PHYSICAL PROOF] [PENDING PREPRESS STANDARD] [PROJECT-SPECIFIC] [REGULATORY REVIEW REQUIRED] [LEGAL REVIEW REQUIRED]")
move(t, top=5.70, height=0.6)
n = clone(s, t, top=6.32, height=0.42, name="Text 25",
          text="Edition states: WORKING EDITION → RELEASE CANDIDATE (RC) → APPROVED PRODUCTION STANDARD. An RC is a synchronized working edition issued for validation; it becomes APPROVED only when all six release gates are met and the Brand Owner authorizes release (AC20).")
set_font_size(n, 11)
L(3, "undeclared labels declared; RC / edition states defined")

# ---------- 4 governance table ----------
s = S(4)
gf = table_of(s)
ren = {"Production Lead": "Production Lead [REQUIRES OWNER] *", "Procurement": "Procurement / Supplier QA", "Localization": "Arabic / Localization Lead",
       "Accessibility": "Accessibility QA", "Legal / Regulatory": "Legal / IP Counsel", "Supplier / Fabricator": "Supplier / Fabricator (external party)"}
rows = gf.table.rows
for r in range(1, len(rows)):
    name = cell(gf, r, 0).text.strip()
    if name in ren: set_cell(cell(gf, r, 0), ren[name])
# merge the separate 'Supplier QA' row into Procurement / Supplier QA
for r in range(1, len(rows)):
    if cell(gf, r, 0).text.strip() == "Supplier QA":
        set_cell(cell(gf, r - 1, 2), "Terms · proof, goods"); del_row(gf, r); break
t = find(s, "Named holders:"); defit(t)
set_text(t, "Named holders: [REQUIRES OWNER]. Roles are the twelve of the V3 Approval Register (Owners & Governance; Part A slide 72). * Production Lead is not a register role: proposed as a delegate of Procurement / Supplier QA, for the Brand Owner to confirm. The allocation is a proposal, not an approved RACI.")
move(t, top=5.45, height=0.6)
set_notes(s, "Role names follow the register's twelve-role list; Part A slide 72 is the canonical statement. The allocation of prepare / check / approve / release / archive is a proposed framework for the Brand Owner to confirm (A06, AC19). No names are supplied.")
L(4, "roles renamed to register roles; Production Lead flagged; Supplier QA merged")

# ---------- 7 master artwork register ----------
s = S(7)
gf = table_of(s)
for r in range(1, len(gf.table.rows)):
    if cell(gf, r, 0).text.strip() == "Standalone Ink symbol":
        set_cell(cell(gf, r, 0), "Standalone Black symbol"); set_cell(cell(gf, r, 1), "PNG supplied")
t = find(s, "PNG is for raster use only"); defit(t)
set_text(t, "PNG is for raster use only. Supplied rasters: GEM_Logo_Horizontal_Beige / Black / White and GEM_Symbol_Beige / Black (Part B); controlled names follow X12 at first manifest (VAL-16). Missing files are never redrawn.")
move(t, top=6.1, height=0.5)
L(7, "Ink symbol row → Black symbol; file list aligned with Part B")

# ---------- 12 / 13 physical reference promotion ----------
t = find(S(12), "Different processes may need"); defit(t)
set_text(t, "Different processes may need different technical values to match the same GEM target. The visual target is the standard; the value is the evidence. The first formally approved physical proof for a process and substrate becomes the controlled process-specific reference target and is logged as a golden sample (Appendix M).")
move(t, height=1.3)
t = find(S(13), "Approval follows printer"); defit(t)
set_text(t, "Approval follows printer or supplier proof validation. Until then no Delta E value, instrument or condition may be quoted to a supplier as a GEM standard. The first approved proof for a process and substrate is promoted to the golden sample register (Appendix M) and becomes the comparison reference for that process. No value is created by this standard.")
move(t, height=0.85)
L(12, "first approved proof → process reference"); L(13, "promotion rule")

# ---------- 14 footer ----------
t = find(S(14), "PART C · COLOUR"); defit(t); set_text(t, "GEM™ PRODUCTION STANDARDS V3.0 · PART C · COLOUR"); move(t, width=6.5)
L(14, "footer prefix restored")

# ---------- 17 notes ----------
set_notes(S(17), "Arabic mechanics and numeral mechanics follow Part B (M11, M12, DS06): guest-facing Arabic may use Arabic-Indic numerals where approved for the locale or property; technical identifiers always remain Latin digits. Part A states the principle; Part B owns the implementation rule. Locale policy is pending sign-off.")
L(17, "numeral ownership note → Part B")

# ---------- 20 safe zones ----------
s = S(20)
t = find(s, "No universal dimensions"); defit(t)
set_text(t, "Safe-zone and bleed roles are standardised (T13); their dimensions come from the supplier dieline, never from this standard.")
move(t, height=0.5)
for n in ["Shape 5", "Text 6", "Shape 7", "Text 8", "Text 9"]: shift(by_name(s, n), 0.22)
L(20, "T13 roles vs dimensions clarified")

# ---------- 27 tiers ----------
s = S(27)
set_text(find(s, "Core Collection"), "Core Collection [REQUIRES OWNER]")
for n in ["Text 4", "Text 7", "Text 10"]: set_text(by_name(s, n), "CONDITIONAL · AC03 / VAL-01")
t = find(s, "It is not a separate brand"); defit(t)
set_text(t, "It is not a separate brand, symbol or colour system. Every pack carries the one GEM identity. All tiers are CONDITIONAL until commercial scope closes (AC03, VAL-01); Property Edition and Bespoke programmes are named in the register (E14), the name Core Collection is not.")
move(t, height=0.55)
for n in ["Text 13", "Text 14"]: shift(by_name(s, n), 0.28)
set_notes(s, "Tiers are carried from Part A slide 60 with the same status: CONDITIONAL under AC03 / VAL-01. No packaging data is invented here.")
L(27, "tier status CONDITIONAL AC03/VAL-01, matching Part A")

# ---------- 29 six levels ----------
s = S(29)
set_text(find(s, "Seven levels"), "Six levels, always in this order.")
set_text(by_name(s, "Text 5"), "3 · Product / function")
remove(by_name(s, "Text 6"))
for n in ["Text 7", "Text 8", "Text 9"]: shift(by_name(s, n), -0.565)
set_text(by_name(s, "Text 7"), "4 · Quantity / size"); set_text(by_name(s, "Text 8"), "5 · Instructions / market information"); set_text(by_name(s, "Text 9"), "6 · SKU / artwork version · [PENDING]")
t = find(s, "Top to bottom"); defit(t); set_text(t, "Top to bottom is the order of visual weight. The layout may flex; the order may not. Six levels as approved (T05); Part A slide 61 states the same.")
move(t, height=0.7)
set_notes(s, "Hierarchy approved in the register (T05) and stated in Part A slide 61. Type sizes per level are [PENDING PRODUCTION MASTER] and set per dieline.")
L(29, "hierarchy aligned to the six T05 levels")

# ---------- 39 sign register ----------
gf = table_of(S(39))
add_row(gf, ["External / building identification", "Names the property from outside", "[DEFERRED]", "Deferred (S02)"])
L(39, "external signage row, Deferred (S02)")

# ---------- 41 bilingual ----------
s = S(41)
set_text(find(s, "Which language leads"), "Arabic leads by default for the approved Saudi/GCC context (S07)")
t = find(s, "Noto Sans Arabic is interim"); defit(t)
set_text(t, "Noto Sans Arabic is interim. Native review is required before any sign is fabricated. Any exception to the Arabic-first default is a documented project-specific deviation (Appendix K), only where regulation or operations require it.")
move(t, height=0.5)
set_notes(s, "Bilingual order follows register decision S07 and Part A slide 65: Arabic first and right, English secondary and left, with equivalent information and legibility. Arabic native QA is an open validation gate (VAL-07). Working Arabic text is never treated as approved copy.")
L(41, "Arabic-first default (S07); exceptions as deviations")

# ---------- 44 templates ----------
t = find(S(44), "Layout follows Part A"); defit(t)
set_text(t, "Layout follows Part A. Print process and stock follow the print and materials standards. Templates are an OPEN DELIVERABLE (AB10, AB11, W01–W10): none exists yet.")
L(44, "template deliverable gate")

# ---------- 56 file naming ----------
s = S(56)
set_text(find(s, "51 · FILE NAMING"), "51 · FILE NAMING · APPROVED · REGISTER X12")
set_text(find(s, "GEM_[Category]"), "GEM_[STREAM]_[ASSET]_[VARIANT]_vX.Y_YYYYMMDD.ext")
n = clone(s, find(s, "Underscores separate"), left=0.90, top=4.05, width=11.54, height=0.7, name="Text 9",
          text="Convention X12 of the V3 Approval Register; the date field is the issue date. Proposed, CONDITIONAL [REQUIRES OWNER]: asset ID GEM-[CATEGORY]-[NNN] and SHA-256 checksums, recorded in the manifest (Appendix N).")
defit(n)
L(56, "X12 naming; asset ID / SHA-256 proposals")

# ---------- 61 appendix B ----------
gf = table_of(S(61))
add_row(gf, ["Artwork use (Y06): GEM artwork is used only for the authorized job; no reuse, distribution or adaptation outside the approved scope · LEGAL REVIEW REQUIRED", "[ENTER]"])
L(61, "Y06 artwork-use clause, LEGAL REVIEW REQUIRED")

# ---------- 63 appendix D ----------
s = S(63); gf = table_of(s)
add_row(gf, ["Proof ID", "", "Samples submitted (count)", ""])
for n in ["Text 2", "Text 3"]: shift(by_name(s, n), 0.36)
L(63, "Proof ID, sample count")

# ---------- 64 appendix E ----------
s = S(64); gf = table_of(s)
add_row(gf, ["Grain direction", "", "Supplier datasheet reference", ""])
for n in ["Text 2", "Text 3"]: shift(by_name(s, n), 0.36)
L(64, "grain direction, datasheet reference")

# ---------- 65 appendix F ----------
gf = table_of(S(65))
for r in range(1, len(gf.table.rows)):
    if cell(gf, r, 0).text.strip() == "Pack format and closure": set_cell(cell(gf, r, 0), "Primary pack: format and closure")
for t in ["Secondary pack", "Market(s)", "Language / bilingual status"]: add_row(gf, [t, "[ENTER]"])
L(65, "primary/secondary pack, market, language")

# ---------- 71 appendix L ----------
s = S(71); gf = table_of(s)
add_row(gf, ["Raised by", "", "Date raised / date closed", ""])
for n in ["Text 2", "Text 3"]: shift(by_name(s, n), 0.36)
L(71, "raised by, dates")

# ---------- 73 appendix N ----------
s = S(73)
t = find(s, "File name: GEM_"); defit(t); set_text(t, "File name: GEM_[STREAM]_[ASSET]_[VARIANT]_vX.Y_YYYYMMDD.ext (register X12) · the date field is the issue date; approval dates live in the manifest."); move(t, height=0.45)
t = find(s, "Rights / ownership status and checksum"); defit(t); set_text(t, "Proposed, CONDITIONAL [REQUIRES OWNER]: asset ID GEM-[CATEGORY]-[NNN] · checksum SHA-256 · Rights / ownership status [PENDING] (VAL-04) · No unreleased file may be listed as APPROVED."); move(t, top=4.72, height=0.45)
L(73, "X12 file name; asset ID / checksum proposals")

# ---------- 75 open gates ----------
s = S(75)
items = ["Production vector logo masters · VAL-02", "Micro / no-spark logo · VAL-02", "Stacked logo · VAL-02", "White symbol · VAL-02",
         "Artwork ownership · VAL-04", "Trademark clearance · VAL-03", "Font build and licence · VAL-05", "CMYK / spot validation · VAL-11, J14",
         "Physical colour standards · VAL-09, J16", "Arabic native QA · VAL-07", "Image rights · VAL-06", "Packaging dielines · VAL-09",
         "Regulatory packaging data · VAL-10", "Packaging physical proof · VAL-09", "Signage site standards · VAL-12, deferred",
         "Named governance holders · AC19", "Final release authorization · AC20"]
for i, t in enumerate(items):
    set_text(by_name(s, f"Text {i+2}"), t)
set_text(by_name(s, "Text 0"), "69 · OPEN VALIDATION GATES · SEVENTEEN OPEN · IDS PER V3 APPROVAL REGISTER")
n = clone(s, by_name(s, "Text 19"), left=0.89, top=4.6, width=11.54, height=0.6, name="Text 21",
          text="Gate IDs follow the register's Validation & Evidence sheet (VAL-01 to VAL-21; VAL-18 is not allocated) and the release checkpoints AC17, AC19 and AC20. Digital gates VAL-13, VAL-14, VAL-20 and the cross-document check VAL-21 are tracked in Part B and the RC2 Open Evidence Register.")
defit(n)
set_notes(s, "All seventeen items are PENDING VALIDATION and map to the register's VAL and AC identifiers; Part A slide 78 lists the same IDs. Each closes only with evidence logged in the asset manifest and change log.")
L(75, "gate items carry register VAL / AC IDs")

# ---------- 76 six gates ----------
s = S(76)
for n, t in {"Text 4": "Open · VAL-02, VAL-03, VAL-04", "Text 7": "Open · AC17, VAL-05, VAL-10", "Text 10": "Open · VAL-09, VAL-11",
             "Text 13": "Open · VAL-09, VAL-15; VAL-12 deferred", "Text 16": "Open · AC19, VAL-16, VAL-21", "Text 19": "Open · AC20"}.items():
    set_text(by_name(s, n), t)
t = find(s, "Until all six are met"); defit(t); set_text(t, "Until all six are met: V3.0 RC2 · WORKING EDITION / PENDING PRODUCTION VALIDATION")
set_notes(s, "This edition is Release Candidate 2, a working edition. No gate is closed; gate IDs follow the register.")
L(76, "gates mapped to register IDs; RC2 edition line")

# ---------- 77 document control (duplicated from 76) ----------
s = S(77)
set_text(by_name(s, "Text 0"), "71 · DOCUMENT CONTROL & CHANGE LOG")
set_text(by_name(s, "Text 1"), "One version. One date. One owner.")
tiles = {"Text 3": "Document ID", "Text 4": "GEM-PS-V3.0-RC2", "Text 6": "Version", "Text 7": "V3.0 RC2 (Release Candidate 2)",
         "Text 9": "Edition state", "Text 10": "WORKING EDITION · NOT RELEASED", "Text 12": "Issue date", "Text 13": "2026-10-06",
         "Text 15": "Owner", "Text 16": "Brand Owner [REQUIRES OWNER]", "Text 18": "Approver", "Text 19": "Brand Owner · AC20 PENDING"}
for n, t in tiles.items(): set_text(by_name(s, n), t)
t = by_name(s, "Text 20"); defit(t); set_text(t, "Release status: RC2 synchronized with Part A and Part B RC2 · evidence gates open · Title: GEM™ Production Standards V3.0 — Part C")
t = by_name(s, "Text 21"); defit(t); set_text(t, "Supersedes: Release Candidate 1 (undated) · Related: GEM-BG-V3.0-RC2 (Part A) · GEM-DDS-V3.0-RC2 (Part B) · V3 Approval Register")
n = clone(s, by_name(s, "Text 22"), left=0.89, top=4.95, width=11.54, height=0.8, name="Text 24",
          text="Change log · RC2, 2026-10-06: authority order and domain ownership (2); edition states and labels (3); register role names (4); master register rows (7); physical reference promotion (12, 13); T13 wording (20); tier status (27); six-level hierarchy T05 (29); external signage deferred S02 (39); Arabic-first default S07 (41); template gate (44); X12 naming (56, N); Y06 clause (B); fields added (D, E, F, L); gate IDs (75, 76). Full record: GEM_V3_RC2_Release_Notes.")
defit(n); set_font_size(n, 11)
set_text(by_name(s, "Text 22"), "GEM™ PRODUCTION STANDARDS V3.0 · PART C · DOCUMENT CONTROL")
set_text(by_name(s, "Text 23"), "77")
set_notes(s, "Document control per register A11 and the change-log requirement (Z24, AB19). Appendix O remains the production change log for assets and standards; this slide records the edition history of Part C itself.")
L(77, "new document-control and change-log slide")

# ---------- 78 closing ----------
s = S(78)
set_text(find(s, "RELEASE CANDIDATE 1"), "RELEASE CANDIDATE 2 · WORKING EDITION / PENDING PRODUCTION VALIDATION")
set_text(by_name(s, "Text 3"), "78")
L(78, "closing: RC2")

# ---------- footers ----------
n = replace_everywhere(prs, "GEM™ PRODUCTION STANDARDS V3.0 · PART C", "GEM™ PRODUCTION STANDARDS V3.0 RC2 · PART C")
L(0, f"footer version string on {n} shapes")

# ---------- alt text ----------
DEC = "Decorative background geometry: cropped GEM ring and quarter forms."
ALT = {(1, "Image 0"): (DEC, True), (1, "Image 1"): ("GEM primary horizontal signature, Beige on Ink, supplied raster, PENDING PRODUCTION MASTER", False),
       (6, "Image 0"): (DEC, True), (18, "Image 0"): (DEC, True), (26, "Image 0"): (DEC, True), (37, "Image 0"): (DEC, True), (46, "Image 0"): (DEC, True), (59, "Image 0"): (DEC, True),
       (10, "Image 0"): ("Primary signature inside the dashed one-u clearspace boundary, schematic only.", False),
       (20, "Image 0"): ("Bleed, trim and safe-area schematic: dashed outer bleed, solid trim, dotted inner brand safe area. Not to scale.", False),
       (31, "Image 0"): ("Dieline schematic showing the roles of trim, bleed, safe area and glue flap. No dimensions.", False),
       (40, "Image 0"): ("Signage artwork zone schematic: symbol zone and text lines. No sizes.", False),
       (78, "Image 0"): (DEC, True), (78, "Image 1"): ("GEM primary horizontal signature, Beige on Ink, supplied raster.", False)}
done = 0
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.shape_type == 13:
            d, dec = ALT[(i, sh.name)]; set_alt(sh, d, dec); done += 1
L(0, f"alt text on {done} pictures")

prs.save(F)
for n, w in log: print(f"C {n:>2}: {w}")
