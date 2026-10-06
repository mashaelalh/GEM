"""Final QA: scan the portfolio deck text (slides + notes) for unsupported claims, invented data and
unlabelled concepts; confirm every slide carries a status; list fonts; write 05_qa/QA_Report.md."""
import re, sys, pathlib, zipfile
from pptx import Presentation

ROOT = pathlib.Path(__file__).resolve().parents[1]
DECK = ROOT / "04_release" / "GEM_Amenities_and_Packaging_Concept_Portfolio_v1.0.pptx"
prs = Presentation(DECK)
rows = []
PROHIBITED = [r"manufactured by GEM", r"sustainabl", r"recyclab", r"biodegrad", r"locally (produced|sourced)", r"premium ingredient", r"\borganic\b", r"\bvegan\b", r"cruelty", r"dermatolog", r"\bcertified\b", r"custom MOQ", r"fast delivery", r"world-class", r"\bbest\b", r"leading", r"luxurious", r"unforgettable", r"\bindulge", r"\bpamper", r"\boasis\b"]
INVENTED = [r"\b\d+\s?(ml|mL|g|oz)\b", r"\b\d+\s?(mm|cm|gsm)\b", r"\bPantone\b", r"\bCMYK\b", r"\bPET\b|\bHDPE\b|\bPP\b", r"\$|SAR\s?\d|€|£", r"\bMOQ\s?\d"]
PALETTE = {"12171D", "BCACA7", "020202", "FFFFFF"}
texts = []
for i, s in enumerate(prs.slides, 1):
    t = "\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
    n = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
    texts.append((i, t, n))
    has_status = bool(re.search(r"CONCEPT|CONDITIONAL|PENDING|REQUIRES", t))
    rows.append((f"Slide {i}: status label present", "PASS" if has_status else "FAIL", ""))
    for pat in PROHIBITED:
        for src, body in (("slide", t), ("notes", n)):
            for m in re.finditer(pat, body, re.I):
                ctx = body[max(0, m.start() - 40): m.end() + 40].replace("\n", " ")
                # allowed when negated / quoted as a prohibition
                ok = re.search(r"(not|no|never|avoid|does not|do not|unless|without|prohibit)[^.]{0,80}" + pat, ctx, re.I) or re.search(pat + r"[^.]{0,60}(not stated|is not|are not|unless|conditional|no )", ctx, re.I)
                rows.append((f"Slide {i} {src}: prohibited term /{pat}/", "PASS (negated)" if ok else "FAIL", ctx))
    for pat in INVENTED:
        for m in re.finditer(pat, t + "\n" + n):
            ctx = (t + "\n" + n)[max(0, m.start() - 30): m.end() + 30].replace("\n", " ")
            rows.append((f"Slide {i}: invented value /{pat}/", "FAIL", ctx))
# images exist and palette
z = zipfile.ZipFile(DECK)
media = [n for n in z.namelist() if n.startswith("ppt/media/")]
rows.append(("Embedded media count", "INFO", str(len(media))))
cols = set()
for n in z.namelist():
    if n.startswith("ppt/slides/slide") and n.endswith(".xml"):
        cols |= set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', z.read(n).decode()))
bad = {c.upper() for c in cols} - PALETTE
rows.append(("Slide XML colours within the four-colour palette", "PASS" if not bad else "FAIL", ", ".join(sorted(bad))))
fonts = set()
for n in z.namelist():
    if n.startswith("ppt/slides/slide") and n.endswith(".xml"):
        fonts |= set(re.findall(r'typeface="([^"]+)"', z.read(n).decode()))
rows.append(("Fonts referenced", "PASS" if fonts <= {"Jost", "Inter"} else "FAIL", ", ".join(sorted(fonts))))
rows.append(("Slide count (target 18–24)", "PASS" if 18 <= len(prs.slides) <= 24 else "FAIL", str(len(prs.slides))))
fails = [r for r in rows if r[1] == "FAIL"]
md = ["# QA Report · GEM_Amenities_and_Packaging_Concept_Portfolio_v1.0", "", f"Automated scan of slide text, notes, colours, fonts and media. **{len(fails)} FAIL.**", "", "| Check | Result | Context |", "|---|---|---|"]
for r in rows: md.append("| " + " | ".join(str(x).replace("|", "/") for x in r) + " |")
md += ["", "## Manual checks (recorded after rendering)", "", "See `Visual_QA.md`: all pages rendered, text overflow, logo distortion, concept labels, family coherence, premium feel, B2B readability.", ""]
(ROOT / "05_qa" / "QA_Report.md").write_text("\n".join(md), encoding="utf-8")
print(len(rows), "checks,", len(fails), "FAIL")
for r in fails: print("  FAIL", r)
