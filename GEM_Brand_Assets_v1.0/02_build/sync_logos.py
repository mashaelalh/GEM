"""Swap the old low-resolution logo rasters inside PPTX decks for the supplied vector kit.
Each picture keeps its box on the slide. The kit SVG is placed, unscaled in proportion, inside the same
visible bounds the old raster had (wrapper viewBox only; no path is edited). The SVG is embedded with a
high-resolution PNG fallback (Office svgBlip extension), so PowerPoint uses vector and LibreOffice uses PNG."""
import io, re, pathlib, hashlib
import numpy as np
from PIL import Image
import cairosvg
from lxml import etree
from pptx.oxml.ns import qn
from pptx.opc.package import Part
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu

KIT = pathlib.Path(__file__).resolve().parents[1] / "04_official_kit" / "logo"
OUTD = pathlib.Path(__file__).resolve().parents[1] / "05_layout_matched"
COLS = {(188, 172, 167): "beige", (2, 2, 2): "black", (255, 255, 255): "white"}
SVG_EXT = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"; ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
_cache = {}

def kit_svg(kind, colour):
    return (KIT / {"horizontal": "horizontal", "symbol": "symbol", "stacked": "stacked", "nospark": "symbol-nospark"}[kind] / "svg" /
            f"gem-{ {'horizontal':'horizontal','symbol':'symbol','stacked':'stacked','nospark':'symbol-nospark'}[kind] }-{colour}.svg").read_text()

def classify(blob):
    im = Image.open(io.BytesIO(blob)).convert("RGBA")
    kind = {(654, 207): "horizontal", (229, 207): "symbol"}.get(im.size)
    if not kind: return None
    a = np.array(im); al = a[..., 3] > 40
    c = tuple(int(round(v)) for v in a[al][:, :3].mean(0))
    colour = min(COLS, key=lambda k: sum((x - y) ** 2 for x, y in zip(k, c)))
    bb = Image.fromarray((al * 255).astype("uint8")).getbbox()
    return kind, COLS[colour], im.size, bb

def matched(kind, colour, size, bbox, out_w_px, name=None):
    """Return (png_bytes, svg_text) with the kit artwork fitted in the old raster's visible bounds."""
    key = (kind, colour, size, bbox, out_w_px)
    if key in _cache: return _cache[key]
    svg = kit_svg(kind, colour)
    m = re.search(r'viewBox="([\d.\- ]+)"', svg); vx, vy, vw, vh = map(float, m.group(1).split())
    W, H = size; bx0, by0, bx1, by1 = bbox; bw, bh = bx1 - bx0, by1 - by0
    k = min(bw / vw, bh / vh); ox = bx0 + (bw - vw * k) / 2; oy = by0 + (bh - vh * k) / 2
    X0, Y0 = vx - ox / k, vy - oy / k
    new_vb = f'{X0:.4f} {Y0:.4f} {W / k:.4f} {H / k:.4f}'
    wrapped = re.sub(r'viewBox="[\d.\- ]+" width="[\d.]+" height="[\d.]+"', f'viewBox="{new_vb}"', svg, count=1)
    assert new_vb in wrapped
    png = cairosvg.svg2png(bytestring=wrapped.encode(), output_width=out_w_px)
    _cache[key] = (png, wrapped)
    if name:
        (OUTD / f"{name}.svg").write_text(wrapped); (OUTD / f"{name}.png").write_bytes(png)
    return _cache[key]

def add_svg_ext(slide, pic, svg_text):
    pkg = slide.part.package
    key = hashlib.sha256(svg_text.encode()).hexdigest()
    store = pkg.__dict__.setdefault("_svg_parts", {})
    if key not in store:
        store[key] = Part(pkg.next_image_partname("svg"), "image/svg+xml", pkg, svg_text.encode())
    rId = slide.part.relate_to(store[key], RT.IMAGE)
    blip = pic._element.blipFill.find(qn("a:blip"))
    for old in blip.findall(qn("a:extLst")): blip.remove(old)
    ext = etree.SubElement(etree.SubElement(blip, qn("a:extLst")), qn("a:ext")); ext.set("uri", SVG_EXT)
    sb = etree.SubElement(ext, "{%s}svgBlip" % ASVG, nsmap={"asvg": ASVG}); sb.set(qn("r:embed"), rId)

def swap_pic(slide, pic, png, svg_text):
    _, rId = slide.part.get_or_add_image_part(io.BytesIO(png))
    pic._element.blipFill.find(qn("a:blip")).set(qn("r:embed"), rId)
    add_svg_ext(slide, pic, svg_text)

def sync_deck(prs, descr_prefix="GEM logo"):
    n = 0
    for s in prs.slides:
        for sh in list(s.shapes):
            if sh.shape_type != 13: continue
            c = classify(sh.image.blob)
            if not c: continue
            kind, colour, size, bb = c
            png, svg = matched(kind, colour, size, bb, size[0] * (6 if kind == "horizontal" else 8), name=f"{kind}-{colour}_{size[0]}x{size[1]}")
            swap_pic(s, sh, png, svg)
            d = sh._element.nvPicPr.cNvPr
            if not d.get("descr"): d.set("descr", f"{descr_prefix} ({'lockup' if kind == 'horizontal' else 'symbol'}, {colour})")
            n += 1
    return n

def add_kit_pic(slide, kind, colour, left, top, width, height, descr, fallback_px=1024):
    """Add a kit logo (SVG + PNG fallback) fitted inside the given box, centred, proportions kept."""
    svg = kit_svg(kind, colour); m = re.search(r'viewBox="([\d.\- ]+)"', svg); vx, vy, vw, vh = map(float, m.group(1).split())
    k = min(width / vw, height / vh); w, h = int(vw * k), int(vh * k)
    png = cairosvg.svg2png(bytestring=svg.encode(), output_width=max(fallback_px, 1))
    pic = slide.shapes.add_picture(io.BytesIO(png), int(left + (width - w) / 2), int(top + (height - h) / 2), w, h)
    pic._element.nvPicPr.cNvPr.set("descr", descr); add_svg_ext(slide, pic, svg)
    return pic
