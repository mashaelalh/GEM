"""Part D RC2 (approved concept portfolio): extract the approved mockups unchanged, swap the low-resolution
logo rasters for the supplied vector kit, and update the logo-status wording. Content, layout and mockups are
otherwise untouched."""
import sys, io, hashlib, pathlib, csv
ROOT = pathlib.Path(__file__).resolve().parents[1]; REPO = ROOT.parent
sys.path.insert(0, str(REPO / "GEM_Brand_Assets_v1.0/02_build")); sys.path.insert(0, str(REPO / "GEM_V3_RC2/05_logs"))
from pptx import Presentation
from PIL import Image
from pptx_helpers import find, set_text
import sync_logos as S
SRC = ROOT / "00_originals" / "GEM Amenities & Packaging — Concept Product Portfolio V3.0 — Part D — RC2 (as supplied).pptx"
OUT = ROOT / "04_release" / "GEM Amenities & Packaging — Concept Product Portfolio V3.0 — Part D — RC2.pptx"
prs = Presentation(SRC); reg = []; seen = {}
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.shape_type != 13: continue
        blob = sh.image.blob; h = hashlib.sha256(blob).hexdigest()
        if S.classify(blob): continue
        im = Image.open(io.BytesIO(blob))
        if h not in seen:
            seen[h] = f"slide-{i:02d}_{im.size[0]}x{im.size[1]}_{h[:8]}.png"
            (ROOT / "01_mockups" / seen[h]).write_bytes(blob)
        reg.append((i, seen[h], im.size[0], im.size[1], h, f"{sh.width/914400:.2f} x {sh.height/914400:.2f} in"))
with open(ROOT / "01_mockups" / "MOCKUP_REGISTER.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["slide", "file", "px_w", "px_h", "sha256", "placed_size"]); w.writerows(reg)
print("mockups extracted:", len(seen), "placements:", len(reg))
print("logos swapped:", S.sync_deck(prs, "GEM logo"))
OLD = "Supplied raster logos remain PENDING PRODUCTION MASTER."
NEW = "Logos are the supplied vector kit (SVG), used unchanged; acceptance as production master remains PENDING (VAL-02)."
n = 0
for s in prs.slides:
    if s.has_notes_slide:
        for p in s.notes_slide.notes_text_frame.paragraphs:
            for r in p.runs:
                if OLD in r.text: r.text = r.text.replace(OLD, NEW); n += 1
                if "Original raster replaces" in r.text: r.text = r.text.replace("Original raster replaces", "Original GEM artwork (supplied vector kit) replaces"); n += 1
print("notes edited:", n)
set_text(find(prs.slides[9], "One supplied raster"), "Supplied vector kit; production master acceptance pending.")
prs.save(OUT); print("saved", OUT.name)
