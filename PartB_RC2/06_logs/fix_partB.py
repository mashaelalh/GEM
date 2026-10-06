import sys, copy, pathlib
sys.path.insert(0, "/home/user/GEM/GEM_V3_RC2/05_logs")
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx_helpers import *
F = pathlib.Path("../01_source/B-RC2.pptx"); prs = Presentation(F); S = lambda n: prs.slides[n-1]
def table_of(s, i=0): return [sh for sh in s.shapes if sh.shape_type == 19][i]
def paragraphs(shape, lines, size=None):
    """One paragraph per line, cloning the first paragraph's properties (keeps auto-numbering)."""
    tf = shape.text_frame; p0 = tf.paragraphs[0]
    set_text(shape, lines[0])
    for ln in lines[1:]:
        el = copy.deepcopy(p0._p); tf._txBody.append(el)
        p = tf.paragraphs[-1]; p.runs[0].text = ln
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    if size:
        for p in tf.paragraphs:
            for r in p.runs: r.font.size = Pt(size)
# slide 4
s = S(4); gf = table_of(s)
set_cell(gf.table.cell(3, 2), "Issued as RC2 · WORKING EDITION / PENDING PRODUCTION VALIDATION (DS04, AC17)")
t = find(s, "1 · V3 Final Brand"); paragraphs(t, ["V3 Final Brand Approval Register (Evidence Required stays conditional; Deferred stays deferred)", "Part A, Brand Guidelines: brand strategy, identity, expression", "Part B, this system: digital implementation", "Part C, Production Standards: physical production", "Formal technical and legal standards where they apply (WCAG 2.2 AA)", "Design System v2.1, Brand Guidelines V2/2.1, V1: historical, superseded", "External benchmarks and inspiration only (RF01–RF03)"], size=11)
move(t, top=4.42, height=2.0)
for n in ["Authority order", "Source of truth"]: move(find(s, n), top=4.08)
move(by_name(s, "Shape 5"), top=4.08)
t = find(s, "Where sources conflict"); set_text(t, "Where sources conflict, the register wins. Domain ownership: where Parts A, B and C cover different implementation domains, the owning document governs that domain unless the register overrides it. Decision IDs (L09, Q17, M05…) are cited so every rule traces back. Futura PT, Montserrat, Aeonik, Satoshi and Whyte are legacy only."); set_font_size(t, 11); move(t, top=4.42, height=1.2)
t = find(s, "Read order:"); set_font_size(t, 11); move(t, top=5.68, height=0.7)
# slide 5
s = S(5)
for n in ["Canonical role list", "Accessibility exceptions"]: move(find(s, n), top=5.0, height=1.1)
# slide 21
s = S(21); gf = table_of(s)
r = len(gf.table.rows) - 1
set_cell(gf.table.cell(r, 1), "[REQUIRES OWNER]"); set_cell(gf.table.cell(r, 2), "Localization Lead; no price invented")
move(find(s, "DatePicker. Props"), top=4.85)
# slide 29
s = S(29); gf = table_of(s)
set_cell(gf.table.cell(2, 1), "Arabic first and right, English secondary and left, equal information: the default for Saudi/GCC (S07, Part A 65). Exceptions only by documented project deviation (Part C 41, Appendix K)")
move(find(s, "Viewing distances"), top=4.95)
# slide 34
s = S(34)
t = find(s, "Document ID · GEM-DDS")
paragraphs(t, ["Document ID · GEM-DDS-V3.0-RC2", "Title · GEM™ Digital Design System V3.0 — Part B", "Version · V3.0 RC2 (Release Candidate 2) · semantic 3.0.0-rc.2", "Edition state · WORKING SPECIFICATION · NOT RELEASED", "Issue date · 2026-10-06", "Owner · Digital Design Lead [REQUIRES OWNER]", "Approver · Brand Owner · AC20 PENDING", "Release status · RC2 synchronized with Parts A and C · evidence gates open", "Supersedes · V3.0 working specification, 3.0.0-draft (undated)", "Related · GEM-BG-V3.0-RC2 (Part A) · GEM-PS-V3.0-RC2 (Part C) · V3 Approval Register"], size=11)
move(t, height=3.2)
t = find(s, "3.0.0-rc.2 · 2026-10-06"); set_font_size(t, 10.5); set_text(t, "3.0.0-rc.2 · 2026-10-06 · RC2 synchronization with Parts A and C and the register: version and document ID; authority order and domain ownership; stale Part C references; register role names; one navigation rule; S07 bilingual default with exception rule; six-level hierarchy cross-reference; family-name policy placeholder; T13 wording; template gates; VAL-18 note; X12 naming; currency placeholder; photography cross-reference; token excerpt label; warning-glyph note; booking-flow test; document control. No token, component or rule value changed."); move(t, top=2.32, height=2.2)
t = find(s, "3.0.0-draft. Updated"); set_font_size(t, 10.5); move(t, top=4.62, height=0.9)
t = find(s, "2.1: initial system"); set_font_size(t, 10.5); move(t, top=5.6)
prs.save(F); print("fixed")
